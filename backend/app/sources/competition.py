"""Loads the competition dataset (data panitia): incident database, weekly condition monitoring,
hourly PI data and the five RCA reports."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from pptx import Presentation

from ..config import N_HOURS, SOURCE_DIR
from ..timeutil import h_of
from .model import Asset, Signal

# Taxonomy harmonisation (ISO 14224 style) for the five RCA cases.
FM_FIX = {("PU-2101B", "Mechanical"): "Leakage", ("KO-3201", "High"): "High Vibration", ("PM-4405B", "Motor"): "Overheat",
          ("HE-3301", "High"): "Fouling", ("BL-5702", "High"): "High Vibration"}
COMP_FIX = {"Journal Bearing": "Bearing", "Motor Bearing": "Bearing"}

CM_KEYS = {
    "PU-2101B": ["vib", "sealflush", "dischp", "btemp"],
    "KO-3201": ["vib", "water", "lop", "bmt"],
    "PM-4405B": ["mtemp", "mvib", "amp", "wtemp"],
    "HE-3301": ["dp", "duty", "cot", "heavy"],
    "BL-5702": ["vib", "h2x", "offset", "btemp"],
}
PI_KEYS = {"FEED": "flow", "DISP": "dischp", "VIB": "vib", "TEMP": "temp", "AMP": "amps", "RATE": "rate"}
# Values stated in the RCA reports plus the planned-intervention assumptions from the analysis report.
META = {
    "PU-2101B": dict(type_code="PU", spare="Pompa cadangan PU-2101A", planned_loss=22.6, planned_h=8,
                     planned_action="Pindah terencana ke PU-2101A, ganti seal cartridge, pulihkan seal-flush >= 6 L/min"),
    "KO-3201": dict(type_code="CO", spare="Tidak ada (Class A, satu train)", planned_loss=594.0, planned_h=12,
                    planned_action="Leak test & plug tube cooler, purifier kontinu, flushing oli saat stop terkendali"),
    "PM-4405B": dict(type_code="EM", spare="Pompa standby PM-4405C", planned_loss=11.2, planned_h=6,
                     planned_action="Jalankan PM-4405C, re-grease/ganti bearing DE motor secara terencana"),
    "HE-3301": dict(type_code="HB", spare="Bypass via spare train", planned_loss=18.4, planned_h=10,
                    planned_action="Bypass ke spare train, hydro-jet cleaning bundle, kendalikan heavy-ends feed"),
    "BL-5702": dict(type_code="BL", spare="Tidak ada (Class A)", planned_loss=205.2, planned_h=6,
                    planned_action="Laser alignment, koreksi soft-foot, ganti elemen kopling saat stop terencana"),
}
INC_COLS = ["id", "mto", "ar", "plant", "tag", "cls", "date", "title", "impact", "prerisk", "score", "pic", "status", "disc",
            "type", "comp", "fm_raw", "dt", "act", "pot", "tot", "due"]


def load_incidents(src: Path = SOURCE_DIR) -> list[dict]:
    df = pd.read_excel(src / "Incident Database" / "Incident Database.xlsx", sheet_name="Incident Database", header=2,
                       keep_default_na=False, na_values=[""])
    out = []
    for r in df.itertuples(index=False):
        row = dict(zip(INC_COLS, [r[0], r[1], r[2], r[3], r[4], r[5], str(r[6])[:10], r[7], r[8], r[9], r[10], r[11], r[12],
                                  r[13], r[14], r[15], r[16], r[17], r[18], r[19], r[20], r[21]]))
        row["id"] = int(row["id"]); row["score"] = int(row["score"])
        for k in ("dt", "act", "pot", "tot"):
            row[k] = float(row[k])
        row["fm"] = FM_FIX.get((row["tag"], row["fm_raw"]), row["fm_raw"])
        row["comp_std"] = COMP_FIX.get(row["comp"], row["comp"])
        row["h"] = h_of(datetime.fromisoformat(row["date"]))
        row["due_h"] = None if row["due"] in (None, "n/a") else h_of(datetime.fromisoformat(str(row["due"])[:10]))
        out.append(row)
    return out


# ------------------------------------------------------------------ RCA reports (tables are drawn as text-box grids)
_HEADERS = {"#", "RC", "Item", "PM No.", "Corrective Action"}


def _texts(slide) -> list[str]:
    return [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def _grid_tables(slide) -> list[list[list[str]]]:
    boxes = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            top, left = round(sh.top / 9525), round(sh.left / 9525)
            if 130 <= top < 680:
                boxes.append((top, left, round(sh.width / 9525), sh.text_frame.text.strip()))
    rows: dict[int, list] = {}
    for top, left, w, txt in boxes:
        key = next((k for k in rows if abs(k - top) <= 4), top)
        rows.setdefault(key, []).append((left, w, txt))
    tables, cur = [], None
    for top in sorted(rows):
        cells = sorted(rows[top])
        if len(cells) == 1 and cells[0][1] >= 1000:
            continue
        texts = [c[2] for c in cells]
        if texts[0] in _HEADERS:
            cur = [texts]; tables.append(cur)
        elif cur is not None:
            cur.append(texts)
    return tables


def _after(texts: list[str], label: str) -> str:
    return next((texts[i + 1] for i, x in enumerate(texts) if x == label and i + 1 < len(texts)), "")


def parse_rca(path: Path) -> dict:
    s = list(Presentation(str(path)).slides)
    t1, t2, t3 = _texts(s[0]), _texts(s[1]), _texts(s[2])
    chrono = [{"t": x, "text": t3[i + 1]} for i, x in enumerate(t3) if re.match(r"^\d{2}-[A-Za-z]{3}-\d{4}", x) and i + 1 < len(t3)]
    p4 = [dict(id=r[0], item=r[1], res=r[2], ev=r[3]) for r in _grid_tables(s[5])[0][1:]]
    m4 = [dict(id=r[0], item=r[1], res=r[2], ev=r[3]) for r in _grid_tables(s[6])[0][1:]]
    capa = []
    for k, tb in enumerate(_grid_tables(s[8])):
        for r in tb[1:]:
            capa.append(dict(rc=r[0], action=r[1], date=r[2], pic=r[3], status=r[4], kind="Korektif" if k == 0 else "Pro-aktif"))
    t10 = _grid_tables(s[9])
    root = next((x.replace("ROOT CAUSE:", "").strip() for x in _texts(s[6]) if x.startswith("ROOT CAUSE:")), "")
    return dict(ar=_after(t1, "AR NUMBER"), pic=_after(t1, "PIC (RCA)"), title=_after(t2, "Title / Problem"),
                date_occ=_after(t2, "Date Occurrence"), date_rep=_after(t2, "Date Reported"), severity=_after(t2, "SEVERITY"),
                immediate=_after(t2, "Immediate Action"), problem=_after(t2, "PROBLEM STATEMENT"), chrono=chrono,
                history=_after(_texts(s[3]), "HISTORICAL DATA & EVIDENCE"), p4=p4, m4=m4, root_cause=root, capa=capa,
                preventive=[dict(rc=r[0], cause=r[1], action=r[2], date=r[3], pic=r[4]) for r in t10[0][1:]],
                pm=[dict(no=r[0], desc=r[1], group=r[2], interval=r[3]) for r in t10[2][1:]])


# ------------------------------------------------------------------ DUMMY fill: one clock for all five assets
# The competition PI data covers only 30 days per asset, a different month for each (Mar, Apr, May, Jun, Jul 2026),
# and weekly condition monitoring covers about six months per asset. Outside those windows the simulator uses a
# labelled DUMMY fill so every asset runs over the whole timeline. The fill only ever shows healthy operation:
# every sign of degradation still comes from the real data.
TRIP_MARGIN_H = 72      # hours before the trip that are never used as examples of normal running


def build_hourly_fill(values: np.ndarray, run: np.ndarray, t0: int, trip_start: int | None, n_hours: int, seed: int) -> np.ndarray:
    """Whole days drawn at random from the asset's own normal days in its real PI window (running, and at least
    TRIP_MARGIN_H before the trip), plus small noise. Day blocks keep the daily pattern."""
    rng = np.random.default_rng(seed)
    end = len(values) if trip_start is None else max(96, min(len(values), trip_start - t0 - TRIP_MARGIN_H))
    first = (-t0) % 24                                   # first 00:00 hour inside the window
    days = [values[s:s + 24] for s in range(first, end - 23, 24) if run[s:s + 24].all() and not np.isnan(values[s:s + 24]).any()]
    pool = np.stack(days)
    normal = values[:end][run[:end]]
    sd = np.nanstd(np.diff(normal, axis=0), axis=0) * 0.25
    fill = pool[rng.integers(0, len(pool), n_hours // 24 + 1)].reshape(-1, values.shape[1])[:n_hours]
    return fill + rng.normal(0.0, 1.0, fill.shape) * sd


def build_weekly_fill(cm_h: np.ndarray, cm_values: np.ndarray, status: list[str], n_hours: int, seed: int):
    """Weekly readings before the first and after the last real reading, on the same weekly grid, scattered around
    the real healthy baseline (mean of the first 4 real readings) with the spread seen in those readings. After the
    real range the asset has been repaired, so it reads as healthy again.
    Returns hours, values, real mask and the real post-repair readings (reused after a planned intervention)."""
    rng = np.random.default_rng(seed)
    first4 = cm_values[:4]
    base = first4.mean(0)
    # Half the real scatter, capped so a DUMMY reading stays within 6% of the baseline: the fill stands for steady
    # healthy running and must never look like drift (the drift rule fires at 10%).
    sd = np.clip(first4.std(0, ddof=1) * 0.5, np.abs(base) * 0.005, np.abs(base) * 0.03)
    trip_i = status.index("TRIP") if "TRIP" in status else None
    post = cm_values[trip_i + 1:] if trip_i is not None and trip_i + 1 < len(cm_values) else first4
    before = np.arange(cm_h[0] - 168, -1, -168)[::-1]
    after = np.arange(cm_h[-1] + 168, n_hours, 168)
    fill = base + np.clip(rng.normal(0.0, 1.0, (len(before) + len(after), cm_values.shape[1])), -2, 2) * sd
    H = np.concatenate([before, cm_h, after]).astype(int)
    V = np.vstack([fill[:len(before)], cm_values, fill[len(before):]])
    real = np.concatenate([np.zeros(len(before), bool), np.ones(len(cm_h), bool), np.zeros(len(after), bool)])
    return H, V, real, post.copy()


def _pi_name(desc: str, tag: str) -> str:
    d = desc.replace(tag, "").strip()
    return d[:1].upper() + d[1:].lower() if d else desc


def load_real_assets(incidents: list[dict], src: Path = SOURCE_DIR) -> list[Asset]:
    assets = []
    for f in sorted((src / "Equipment Performance").glob("*.xlsx")):
        tag = re.search(r"RCA\d (\S+)\.xlsx", f.name).group(1)
        wb = openpyxl.load_workbook(f, data_only=True)
        info = {r[0]: r[1] for r in wb["Equipment Info"].iter_rows(min_row=4, max_row=16, values_only=True) if r[0]}
        weekly = []
        for k, r in zip(CM_KEYS[tag], wb["Equipment Info"].iter_rows(min_row=5, max_row=8, values_only=True)):
            m = re.match(r"^(.*?)\s*\(([^)]*)\)\s*$", r[2])
            a, t = [float(x) for x in str(r[3]).split("/")]
            weekly.append(Signal(k, m.group(1) if m else r[2], m.group(2) if m else "", a, t, 1 if t > a else -1, False))
        perf = {r[0]: r[1] for r in wb["Performance Summary"].iter_rows(min_row=4, values_only=True) if r[0]}
        cm = pd.read_excel(f, sheet_name="Condition History")
        pf = next((src / "Production Data").glob(f"*{tag}*.xlsx"))
        tags = pd.read_excel(pf, sheet_name=0)
        pi = pd.read_excel(pf, sheet_name=1)
        pi["Timestamp"] = pd.to_datetime(pi.Timestamp)
        hourly, cols = [], []
        for r in tags.itertuples(index=False):
            if r.Name == "RUN_STATUS":
                continue
            suffix = r.Name.split("_")[-1]
            key = PI_KEYS.get(suffix, suffix.lower())
            hourly.append(Signal(key, _pi_name(r.Description, tag), str(r.engunits), None, None, 1, r.Name != "PLANT_RATE"))
            cols.append(r.Name)
        off = pi[pi.RUN_STATUS == "OFF"].Timestamp
        rca = parse_rca(next((src / "RCA - Downtime Data").glob(f"*{tag}*.pptx")))
        rca["date_rep_h"] = h_of(datetime.strptime(rca["date_rep"], "%d %b %Y"))
        inc = next(x for x in incidents if x["tag"] == tag)
        loss, dth = float(perf["Estimated Loss (k USD)"]), float(perf["Total Downtime (hours)"])
        name = info["Equipment Name"]
        meta = META[tag]
        plant = info["Plant / Unit"]
        seed = sum(map(ord, tag))
        hr_t0 = h_of(pi.Timestamp.iloc[0].to_pydatetime())
        hr_values, hr_run = pi[cols].to_numpy(dtype=float), (pi.RUN_STATUS == "ON").to_numpy()
        trip_start = h_of(off.min().to_pydatetime())
        cm_h0 = np.array([h_of(pd.Timestamp(d).to_pydatetime()) for d in cm["Date"]])
        cm_h, cm_values, cm_real, cm_post = build_weekly_fill(cm_h0, cm.iloc[:, 2:6].to_numpy(dtype=float), list(cm["Health Status"]),
                                                              N_HOURS, seed + 1)
        assets.append(Asset(
            id=tag, name=name, short=re.sub(r"\(.*?\)", "", name.replace(tag, "")).strip(), eq_type=info["Equipment Type"],
            type_code=meta["type_code"], cls=info["Equipment Class"], plant_code=re.search(r"\((\w+)\)", plant).group(1), plant=plant,
            disc=info["Discipline"], crit=info["Criticality"], pic=rca["pic"], spare=meta["spare"], source="real",
            loss_per_h=round(loss / dth, 2), trip_hours=dth, trip_loss=loss, prod_loss=float(perf["Production Loss (ton)"]),
            planned_loss=meta["planned_loss"], planned_h=meta["planned_h"], planned_action=meta["planned_action"],
            hourly=hourly, weekly=weekly,
            hr_t0=hr_t0, hr_values=hr_values, hr_run=hr_run,
            hr_fill=build_hourly_fill(hr_values, hr_run, hr_t0, trip_start, N_HOURS, seed),
            cm_h=cm_h, cm_values=cm_values, cm_real=cm_real, cm_post=cm_post,
            trip_start=trip_start, trip_end=h_of(off.max().to_pydatetime()) + 1,
            scenario="Kejadian nyata dari data panitia", rca=rca, incident=inc,
            extra={"design_life": info["Design Life"], "monitoring": info["Monitoring Method"],
                   "pm_compliance": float(perf["PM Compliance (%)"]), "mtbf": float(perf["MTBF (hours)"])},
        ))
    return assets
